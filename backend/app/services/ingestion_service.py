import os
import re
import hashlib
import tempfile
from typing import Dict, Any, Optional
from app.services.llm_service import detect_category
from langchain_huggingface import HuggingFaceEmbeddings
from pgvector.psycopg2 import register_vector

import requests
from bs4 import BeautifulSoup

try:
    import fitz  # PyMuPDF
    PYMUPDF_AVAILABLE = True
except ImportError:
    PYMUPDF_AVAILABLE = False

try:
    import pytesseract
    from PIL import Image
    import io
    # Auto-detect tesseract in Docker, fallback to Windows path for local dev
    import shutil
    tesseract_path = shutil.which('tesseract')
    if tesseract_path:
        pytesseract.pytesseract.tesseract_cmd = tesseract_path
    elif os.path.exists(r'C:\Users\Venkata.Jakkinapalli\AppData\Local\Programs\Tesseract-OCR\tesseract.exe'):
        pytesseract.pytesseract.tesseract_cmd = r'C:\Users\Venkata.Jakkinapalli\AppData\Local\Programs\Tesseract-OCR\tesseract.exe'
    OCR_AVAILABLE = True
except ImportError:
    OCR_AVAILABLE = False

try:
    from pptx import Presentation
    PPTX_AVAILABLE = True
except ImportError:
    PPTX_AVAILABLE = False

from app.db.database import get_connection


# ==========================================================
# DEFAULT CONFIG
# ==========================================================

DEFAULT_INGEST_CONFIG = {
    "embedding_model": "sentence-transformers/all-MiniLM-L6-v2",
    "chunk_size": 500,
    "chunk_overlap": 50,
    "kb_files_table": "kb_files",
    "kb_chunks_table": "kb_chunks"
}


# ==========================================================
# UTILS
# ==========================================================

def generate_hash(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def create_embeddings(model_name: str):
    return HuggingFaceEmbeddings(model_name=model_name)


# ==========================================================
# EXTRACTION FUNCTIONS
# ==========================================================

def extract_text_file(path: str) -> str:
    with open(path, "r", encoding="utf-8") as f:
        return f.read()


def extract_html_structured(html: str, base_url: str = "") -> str:
    soup = BeautifulSoup(html, "html.parser")

    # Remove nav, footer, script, style noise
    for tag in soup.find_all(["nav", "footer", "script", "style", "noscript"]):
        tag.decompose()

    content = []
    current_section = ""

    for tag in soup.find_all(["h1", "h2", "h3", "h4", "p", "li", "td", "th", "img"]):
        if tag.name == "img":
            # Try alt text first
            alt = tag.get("alt", "").strip()
            if alt:
                content.append(f"{current_section}: [Image: {alt}]" if current_section else f"[Image: {alt}]")

            # OCR the image if it has a src and OCR is available
            if OCR_AVAILABLE:
                src = tag.get("src", "")
                if src:
                    try:
                        if src.startswith("http"):
                            img_response = requests.get(src, timeout=5)
                            img = Image.open(io.BytesIO(img_response.content))
                        elif src.startswith("data:image"):
                            # base64 embedded image
                            import base64
                            header, data = src.split(",", 1)
                            img = Image.open(io.BytesIO(base64.b64decode(data)))
                        else:
                            continue
                        ocr_text = pytesseract.image_to_string(img).strip()
                        if ocr_text:
                            content.append(f"{current_section}: {ocr_text}" if current_section else ocr_text)
                    except Exception:
                        pass  # skip unreadable images silently
            continue

        text = tag.get_text(strip=True)
        if not text:
            continue

        if tag.name in ["h1", "h2", "h3", "h4"]:
            current_section = text
        else:
            content.append(f"{current_section}: {text}" if current_section else text)

    return "\n".join(content)


def extract_pdf(path: str) -> str:
    if not PYMUPDF_AVAILABLE:
        raise ImportError("PyMuPDF (fitz) is not installed. Run: pip install pymupdf")
    doc = fitz.open(path)
    text = ""
    for page in doc:
        text += page.get_text()
        # OCR images embedded in PDF pages
        if OCR_AVAILABLE:
            for img in page.get_images(full=True):
                try:
                    xref = img[0]
                    base_image = doc.extract_image(xref)
                    img_obj = Image.open(io.BytesIO(base_image["image"]))
                    ocr_text = pytesseract.image_to_string(img_obj).strip()
                    if ocr_text:
                        text += "\n" + ocr_text
                except Exception:
                    pass
    return text


def extract_image_text(path: str) -> str:
    if not OCR_AVAILABLE:
        raise ImportError("pytesseract or Pillow not installed. Run: pip install pytesseract pillow")
    img = Image.open(path)
    return pytesseract.image_to_string(img)


def extract_pptx(path: str) -> str:
    if not PPTX_AVAILABLE:
        raise ImportError("python-pptx not installed. Run: pip install python-pptx")
    prs = Presentation(path)
    lines = []
    for slide_num, slide in enumerate(prs.slides, 1):
        lines.append(f"Slide {slide_num}")
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = para.text.strip()
                    if text:
                        lines.append(text)
    return "\n".join(lines)


def fetch_url_content(url: str) -> str:
    headers = {"User-Agent": "Mozilla/5.0"}
    response = requests.get(url, timeout=15, headers=headers)
    response.raise_for_status()
    return extract_html_structured(response.text, base_url=url)


# ==========================================================
# CLEANING LAYER
# ==========================================================

def clean_text(text: str) -> str:
    text = re.sub(r'\S+@\S+', '[EMAIL]', text)
    text = re.sub(r'http\S+', '[URL]', text)
    text = re.sub(r'password\s*=\s*\S+', '[REDACTED]', text)
    text = re.sub(r'api[_-]?key\s*=\s*\S+', '[REDACTED]', text)
    text = re.sub(r'\s+', ' ', text)
    return text.strip()


# ==========================================================
# SMART CHUNKING
# ==========================================================

def smart_chunk(text: str, max_size: int = 500) -> list:
    # First, normalize whitespace and split into sentences
    text = re.sub(r'\s+', ' ', text).strip()
    
    # Split on sentence boundaries (., !, ?) followed by space
    sentences = re.split(r'(?<=[.!?])\s+', text)
    
    chunks = []
    current = ""
    
    for sentence in sentences:
        sentence = sentence.strip()
        if not sentence:
            continue
        
        # If adding this sentence exceeds max_size, save current and start new
        if current and len(current) + len(sentence) + 1 > max_size:
            chunks.append(current)
            current = sentence
        else:
            current = current + " " + sentence if current else sentence
    
    # Handle leftover
    if current:
        chunks.append(current)
    
    # If no sentence boundaries found (e.g., list items), fall back to word-based chunking
    if len(chunks) == 1 and len(chunks[0]) > max_size:
        words = chunks[0].split()
        chunks = []
        current = ""
        for word in words:
            if len(current) + len(word) + 1 > max_size:
                if current:
                    chunks.append(current)
                current = word
            else:
                current = current + " " + word if current else word
        if current:
            chunks.append(current)
    
    return chunks
# ==========================================================
# MAIN INGEST FUNCTION
# ==========================================================

def ingest_file(
    file_path: str,
    bot_id: int,
    ingest_config: Optional[Dict[str, Any]] = None
) -> dict:

    config = {**DEFAULT_INGEST_CONFIG, **(ingest_config or {})}

    file_name = os.path.basename(file_path)
    ext = file_path.lower().split('.')[-1]

    # ------------------------------------------------------
    # 1️⃣ Extract content based on file type
    # ------------------------------------------------------

    if ext in ["txt", "md"]:
        raw_text = extract_text_file(file_path)
    elif ext == "html":
        with open(file_path, "r", encoding="utf-8") as f:
            raw_text = extract_html_structured(f.read())
    elif ext == "pdf":
        raw_text = extract_pdf(file_path)
    elif ext in ["png", "jpg", "jpeg"]:
        raw_text = extract_image_text(file_path)
    elif ext == "pptx":
        raw_text = extract_pptx(file_path)
    else:
        raise ValueError(f"Unsupported file type: .{ext}")

    if not raw_text or not raw_text.strip():
        return {
            "file_name": file_name,
            "message": "Empty content extracted, skipping file",
            "chunks_inserted": 0,
            "chunks_skipped": 0
        }

    # ------------------------------------------------------
    # 2️⃣ Clean + smart chunk
    # ------------------------------------------------------

    cleaned = clean_text(raw_text)
    chunk_texts = smart_chunk(cleaned, max_size=config["chunk_size"])

    file_hash = generate_hash(cleaned)

    # ------------------------------------------------------
    # 3️⃣ Create Embedding Model (Configurable)
    # ------------------------------------------------------

    embeddings = create_embeddings(config["embedding_model"])

    # ------------------------------------------------------
    # 4️⃣ DB Connection
    # ------------------------------------------------------

    conn = get_connection()
    register_vector(conn)
    cur = conn.cursor()

    inserted_chunks = 0
    skipped_chunks = 0

    kb_files_table = config["kb_files_table"]
    kb_chunks_table = config["kb_chunks_table"]

    try:

        # --------------------------------------------------
        # 5️⃣ Prevent Duplicate File (by hash + bot_id)
        # --------------------------------------------------

        cur.execute(
            f"""
            SELECT id FROM {kb_files_table}
            WHERE file_hash = %s AND bot_id = %s;
            """,
            (file_hash, bot_id)
        )

        existing_file = cur.fetchone()

        if existing_file:
            return {
                "file_name": file_name,
                "message": "File already ingested",
                "chunks_inserted": 0,
                "chunks_skipped": 0
            }

        # --------------------------------------------------
        # 6️⃣ Insert File Record
        # --------------------------------------------------

        cur.execute(
            f"""
            INSERT INTO {kb_files_table}
            (file_name, file_hash, bot_id)
            VALUES (%s, %s, %s)
            RETURNING id;
            """,
            (file_name, file_hash, bot_id)
        )

        kb_file_id = cur.fetchone()[0]

        # --------------------------------------------------
        # 7️⃣ Insert Chunks
        # --------------------------------------------------

        for chunk_text in chunk_texts:
            chunk_text = chunk_text.strip()

            if not chunk_text:
                continue

            chunk_hash = generate_hash(chunk_text)

            # Duplicate check per bot
            cur.execute(
                f"""
                SELECT id FROM {kb_chunks_table}
                WHERE chunk_hash = %s AND bot_id = %s;
                """,
                (chunk_hash, bot_id)
            )

            if cur.fetchone():
                skipped_chunks += 1
                continue

            vector = embeddings.embed_query(chunk_text)

            chunk_category = detect_category(chunk_text)

            cur.execute(
                f"""
                INSERT INTO {kb_chunks_table}
                (chunk_text, chunk_hash, embedding, kb_file_id, bot_id, category)
                VALUES (%s, %s, %s, %s, %s, %s);
                """,
                (chunk_text, chunk_hash, vector, kb_file_id, bot_id, chunk_category)
            )

            inserted_chunks += 1

        conn.commit()

        return {
            "file_name": file_name,
            "chunks_inserted": inserted_chunks,
            "chunks_skipped": skipped_chunks
        }

    except Exception as e:
        conn.rollback()
        raise e

    finally:
        cur.close()
        conn.close()


# ==========================================================
# URL INGESTION (same pipeline, no file needed)
# ==========================================================

def ingest_from_url(
    url: str,
    bot_id: int,
    ingest_config: Optional[Dict[str, Any]] = None
) -> dict:

    config = {**DEFAULT_INGEST_CONFIG, **(ingest_config or {})}

    raw_text = fetch_url_content(url)

    if not raw_text or not raw_text.strip():
        return {
            "source": url,
            "message": "No content extracted from URL",
            "chunks_inserted": 0,
            "chunks_skipped": 0
        }

    cleaned = clean_text(raw_text)
    chunk_texts = smart_chunk(cleaned, max_size=config["chunk_size"])
    file_hash = generate_hash(cleaned)

    # Use domain+path as the "file name" for display
    from urllib.parse import urlparse
    parsed = urlparse(url)
    source_name = parsed.netloc + parsed.path

    embeddings = create_embeddings(config["embedding_model"])

    conn = get_connection()
    register_vector(conn)
    cur = conn.cursor()

    inserted_chunks = 0
    skipped_chunks = 0

    kb_files_table = config["kb_files_table"]
    kb_chunks_table = config["kb_chunks_table"]

    try:
        cur.execute(
            f"SELECT id FROM {kb_files_table} WHERE file_hash = %s AND bot_id = %s;",
            (file_hash, bot_id)
        )
        if cur.fetchone():
            return {
                "source": url,
                "message": "URL already ingested (content unchanged)",
                "chunks_inserted": 0,
                "chunks_skipped": 0
            }

        cur.execute(
            f"INSERT INTO {kb_files_table} (file_name, file_hash, bot_id) VALUES (%s, %s, %s) RETURNING id;",
            (source_name, file_hash, bot_id)
        )
        kb_file_id = cur.fetchone()[0]

        for chunk_text in chunk_texts:
            chunk_text = chunk_text.strip()
            if not chunk_text:
                continue

            chunk_hash = generate_hash(chunk_text)
            cur.execute(
                f"SELECT id FROM {kb_chunks_table} WHERE chunk_hash = %s AND bot_id = %s;",
                (chunk_hash, bot_id)
            )
            if cur.fetchone():
                skipped_chunks += 1
                continue

            vector = embeddings.embed_query(chunk_text)
            chunk_category = detect_category(chunk_text)

            cur.execute(
                f"""INSERT INTO {kb_chunks_table}
                (chunk_text, chunk_hash, embedding, kb_file_id, bot_id, category)
                VALUES (%s, %s, %s, %s, %s, %s);""",
                (chunk_text, chunk_hash, vector, kb_file_id, bot_id, chunk_category)
            )
            inserted_chunks += 1

        conn.commit()
        return {
            "source": url,
            "chunks_inserted": inserted_chunks,
            "chunks_skipped": skipped_chunks
        }

    except Exception as e:
        conn.rollback()
        raise e

    finally:
        cur.close()
        conn.close()